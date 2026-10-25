"""Генерирует бланки .docx с метками {{ ... }} для docxtpl.
Свои бланки можно делать прямо в Word: просто впишите {{ key }} там, где нужны данные.
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).parent / "templates"


def new_doc() -> Document:
    d = Document()
    for s in d.sections:
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)
        s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)
    return d


def p(d, text="", align=None, bold=False, indent=None, first=None, after=0):
    para = d.add_paragraph()
    if text:
        para.add_run(text).bold = bold
    if align is not None:
        para.alignment = align
    if indent is not None:
        para.paragraph_format.left_indent = Cm(indent)
    if first is not None:
        para.paragraph_format.first_line_indent = Cm(first)
    para.paragraph_format.space_after = Pt(after)
    return para


def memo():
    d = new_doc()
    p(d, "{{ from_dept }}", bold=True, after=14)
    p(d, "Кому: {{ to_whom }}", indent=8.5, after=24)
    p(d, "СЛУЖЕБНАЯ ЗАПИСКА", WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=6)
    p(d, "{{ subject }}", WD_ALIGN_PARAGRAPH.CENTER, after=18)
    p(d, "{{ body }}", WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.25, after=12)
    p(d, "Срок ответа или решения: {{ deadline }}", after=36)
    p(d, "{{ author_post }}          ______________ / {{ author_fio }}")
    p(d, "{{ today }}")
    d.save(OUT / "memo.docx")


def order():
    d = new_doc()
    p(d, "ПОРУЧЕНИЕ (проект)", WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=6)
    p(d, "от {{ today }}", WD_ALIGN_PARAGRAPH.CENTER, after=18)
    p(d, "Исполнитель: {{ executor_fio }}, {{ executor_dept }}", after=10)
    p(d, "Содержание задания: {{ task }}", WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
    p(d, "Срок исполнения: {{ deadline }}", after=10)
    p(d, "Контроль исполнения возложить на: {{ controller_fio }}", after=36)
    p(d, "Визы согласующих: ______________________")
    d.save(OUT / "order.docx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    memo()
    order()
    print("Бланки созданы в", OUT)
